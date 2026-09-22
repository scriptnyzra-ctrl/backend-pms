import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

BASE_URL = "https://driver-api.rideblitz.id"
HEADERS_BASE = {
    "accept": "application/json, text/plain, */*",
    "accept-language": "en",
    "cache-control": "no-cache",
    "content-type": "application/json",
    "origin": "https://admin-manage.rideblitz.id",
    "pragma": "no-cache",
    "referer": "https://admin-manage.rideblitz.id/",
    "sec-ch-ua": '"Chromium";v="148", "Google Chrome";v="148", "Not/A)Brand";v="99"',
    "sec-ch-ua-mobile": "?0",
    "sec-ch-ua-platform": '"Windows"',
    "sec-fetch-dest": "empty",
    "sec-fetch-mode": "cors",
    "sec-fetch-site": "same-site",
    "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36",
}


def login(username: str, password: str) -> str:
    resp = requests.post(
        f"{BASE_URL}/panel/login",
        headers=HEADERS_BASE,
        json={"username": username, "password": password},
        timeout=30,
    )
    resp.raise_for_status()
    body = resp.json()
    if not body.get("result"):
        raise ValueError(f"Login gagal: {body}")
    token = body["data"]["access_token"]
    print(f"[OK] Login berhasil sebagai '{body['data']['username']}'")
    return token


def fetch_driver_list(token: str, page: int = 1, offset: int = 100) -> dict:
    params = [
        ("sort", "-1"), ("status", "1"), ("status", "2"), ("status", "8"),
        ("status", "3"), ("status", "4"), ("status", "5"), ("status", "6"),
        ("status", "7"), ("attendance", ""), ("page", page), ("offset", offset),
        ("term", ""), ("app_version_name", ""), ("bank_info_provided", "undefined"),
    ]
    headers = {**HEADERS_BASE, "authorization": token}
    resp = requests.get(
        f"{BASE_URL}/v2/panel/driver-list",
        headers=headers,
        params=params,
        timeout=30,
    )
    resp.raise_for_status()
    body = resp.json()
    if not body.get("result"):
        raise ValueError(f"Gagal mengambil driver list: {body}")
    return body["data"]


def fetch_all_drivers(token: str, offset: int = 100) -> list:
    first_page = fetch_driver_list(token, page=1, offset=offset)
    total_pages = first_page["total_pages"]
    total_records = first_page["total_records"]
    print(f"[OK] Total driver: {total_records}, halaman: {total_pages}")

    all_drivers = list(first_page["driver_list_response"])

    for page in range(2, total_pages + 1):
        data = fetch_driver_list(token, page=page, offset=offset)
        all_drivers.extend(data["driver_list_response"])
        print(f"  Mengambil halaman {page}/{total_pages}...", end="\r")

    print(f"\n[OK] Total data diambil: {len(all_drivers)} driver")
    return all_drivers


def flatten_driver(record: dict) -> dict:
    d = record.get("drivers", {})
    vendor = d.get("vendor", {})
    account_state = d.get("account_state", {})
    return {
        "id": d.get("id"),
        "name": d.get("name"),
        "phone_prefix": d.get("phone_prefix"),
        "phone_number": d.get("phone_number"),
        "city_id": d.get("city_id"),
        "city_name": d.get("city_name"),
        "account_state_id": d.get("account_state_id"),
        "account_state_status": account_state.get("status"),
        "otp": d.get("otp"),
        "reason": d.get("reason"),
        "token": d.get("token"),
        "android_version": d.get("android_version"),
        "app_version": d.get("app_version"),
        "app_version_name": d.get("app_version_name"),
        "app_android_version": d.get("app_android_version"),
        "attendance_status": d.get("attendance_status"),
        "last_active": d.get("last_active"),
        "gps_status": d.get("gps_status"),
        "gps_status_uat": d.get("gps_status_uat"),
        "gps_spoof_status": d.get("gps_spoof_status"),
        "gps_spoof_status_uat": d.get("gps_spoof_status_uat"),
        "device_integrity_status": d.get("device_integrity_status"),
        "user_id": d.get("user_id"),
        "bank_info_provided_driver": d.get("bank_info_provided"),
        "vendor_id": d.get("vendor_id"),
        "vendor_name": vendor.get("name"),
        "vendor_description": vendor.get("description"),
        "vendor_is_active": vendor.get("is_active"),
        "assignment_allowed": d.get("assignment_allowed"),
        "rating": d.get("rating"),
        "created_at": d.get("created_at"),
        "updated_at": d.get("updated_at"),
        "hubs": record.get("hubs"),
        "businesses": record.get("businesses"),
        "registered_at": record.get("registered_at"),
        "bank_info_provided": record.get("bank_info_provided"),
    }


def save_to_excel(drivers: list, filename: str):
    if not drivers:
        print("[WARN] Tidak ada data driver untuk disimpan.")
        return

    rows = [flatten_driver(r) for r in drivers]
    columns = list(rows[0].keys())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Driver List"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", start_color="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").upper())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    alt_fill = PatternFill("solid", start_color="D6E4F0")
    data_font = Font(name="Arial", size=10)
    data_align = Alignment(vertical="center")

    for row_idx, row in enumerate(rows, start=2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, key in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[key])
            cell.font = data_font
            cell.alignment = data_align
            if fill:
                cell.fill = fill

    ws.row_dimensions[1].height = 30
    for col_idx, col_name in enumerate(columns, start=1):
        values = [len(col_name)] + [len(str(r[col_name])) for r in rows if r[col_name] is not None]
        max_len = max(values) if values else 10
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"[OK] Data disimpan ke '{filename}' ({len(rows)} baris, {len(columns)} kolom)")


def main():
    username = "merapi"
    password = "qRKzNbami4"

    try:
        token = login(username, password)
    except Exception as e:
        print(f"[ERROR] Login gagal: {e}")
        return

    try:
        drivers = fetch_all_drivers(token, offset=100)
    except Exception as e:
        print(f"[ERROR] Gagal mengambil data driver: {e}")
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = f"driver_list_{timestamp}.xlsx"
    save_to_excel(drivers, output_file)


if __name__ == "__main__":
    main()