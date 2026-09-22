import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading

BASE_URL = "https://driver-api.rideblitz.id"
USER_URL = "https://user.rideblitz.id"
HEADERS_BASE = {
    "accept": "application/json, text/plain, */*",
    "accept-language": "en",
    "cache-control": "no-cache",
    "content-type": "application/json",
    "origin": "https://admin-manage.rideblitz.id",
    "pragma": "no-cache",
    "referer": "https://admin-manage.rideblitz.id/",
    "sec-ch-ua": '"Chromium";v="148", "Google Chrome";v="148", "Not/A)Brand";v="99"',
    "sec-ch-ua-mobile": "?0",
    "sec-ch-ua-platform": '"Windows"',
    "sec-fetch-dest": "empty",
    "sec-fetch-mode": "cors",
    "sec-fetch-site": "same-site",
    "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36",
}

_print_lock = threading.Lock()


def safe_print(msg):
    with _print_lock:
        print(msg)


def login(username: str, password: str) -> str:
    resp = requests.post(
        f"{BASE_URL}/panel/login",
        headers=HEADERS_BASE,
        json={"username": username, "password": password},
        timeout=30,
    )
    resp.raise_for_status()
    body = resp.json()
    if not body.get("result"):
        raise ValueError(f"Login gagal: {body}")
    token = body["data"]["access_token"]
    print(f"[OK] Login berhasil sebagai '{body['data']['username']}'")
    return token


def fetch_driver_list(token: str, page: int = 1, offset: int = 100) -> dict:
    params = [
        ("sort", "-1"), ("status", "1"), ("status", "2"), ("status", "8"),
        ("status", "3"), ("status", "4"), ("status", "5"), ("status", "6"),
        ("status", "7"), ("attendance", ""), ("page", page), ("offset", offset),
        ("term", ""), ("app_version_name", ""), ("bank_info_provided", "undefined"),
    ]
    headers = {**HEADERS_BASE, "authorization": token}
    resp = requests.get(
        f"{BASE_URL}/v2/panel/driver-list",
        headers=headers,
        params=params,
        timeout=30,
    )
    resp.raise_for_status()
    body = resp.json()
    if not body.get("result"):
        raise ValueError(f"Gagal mengambil driver list: {body}")
    return body["data"]


def fetch_all_drivers(token: str, offset: int = 100) -> list:
    first_page = fetch_driver_list(token, page=1, offset=offset)
    total_pages = first_page["total_pages"]
    total_records = first_page["total_records"]
    print(f"[OK] Total driver: {total_records}, halaman: {total_pages}")

    all_drivers = list(first_page["driver_list_response"])
    for page in range(2, total_pages + 1):
        data = fetch_driver_list(token, page=page, offset=offset)
        all_drivers.extend(data["driver_list_response"])
        print(f"  Mengambil halaman {page}/{total_pages}...", end="\r")

    print(f"\n[OK] Total data diambil: {len(all_drivers)} driver")
    return all_drivers


def fetch_driver_profile(token: str, driver_id: int) -> dict:
    headers = {**HEADERS_BASE, "authorization": token}
    try:
        resp = requests.get(
            f"{BASE_URL}/panel/driver-profile/{driver_id}",
            headers=headers,
            timeout=30,
        )
        resp.raise_for_status()
        body = resp.json()
        return body.get("data", {}) if body.get("result") else {}
    except Exception:
        return {}


def fetch_bank_detail(token: str, driver_id: int) -> dict:
    headers = {**HEADERS_BASE, "authorization": token}
    try:
        resp = requests.get(
            f"{USER_URL}/v1/app/users/bank_detail/drivers/{driver_id}",
            headers=headers,
            timeout=30,
        )
        resp.raise_for_status()
        body = resp.json()
        return body.get("data", {}) if body.get("result") else {}
    except Exception:
        return {}


def fetch_driver_detail(token: str, driver_id: int) -> tuple:
    profile = fetch_driver_profile(token, driver_id)
    bank = fetch_bank_detail(token, driver_id)
    return driver_id, profile, bank


def flatten_driver(record: dict, profile: dict, bank: dict) -> dict:
    d = record.get("drivers", {})
    vendor = d.get("vendor", {})
    account_state = d.get("account_state", {})

    dp = profile.get("driver_profile", {})
    dd = dp.get("driver_Details", {})
    attendance_latest = dp.get("attendance", {})
    coords = profile.get("current_cordinates", {})
    documents = dp.get("documents", [])
    bank_profile = dp.get("bank_details", {})

    doc_selfie = next((x for x in documents if x.get("fields", {}).get("key") == "selfie"), {})
    doc_ktp = next((x for x in documents if x.get("fields", {}).get("key") == "ktp"), {})
    doc_sim = next((x for x in documents if x.get("fields", {}).get("key") == "sim"), {})

    driver_businesses = profile.get("driver_businesses", [])
    driver_hubs = profile.get("driver_hubs", [])

    return {
        "id": d.get("id"),
        "name": d.get("name"),
        "phone_prefix": d.get("phone_prefix"),
        "phone_number": d.get("phone_number"),
        "city_id": d.get("city_id"),
        "city_name": d.get("city_name"),
        "account_state_id": d.get("account_state_id"),
        "account_state_status": account_state.get("status"),
        "otp": d.get("otp"),
        "reason": d.get("reason"),
        "token": d.get("token"),
        "android_version": d.get("android_version"),
        "app_version": d.get("app_version"),
        "app_version_name": d.get("app_version_name"),
        "app_android_version": d.get("app_android_version"),
        "attendance_status": d.get("attendance_status"),
        "last_active": d.get("last_active"),
        "gps_status": d.get("gps_status"),
        "gps_status_uat": d.get("gps_status_uat"),
        "gps_spoof_status": d.get("gps_spoof_status"),
        "gps_spoof_status_uat": d.get("gps_spoof_status_uat"),
        "device_integrity_status": d.get("device_integrity_status"),
        "user_id": d.get("user_id"),
        "bank_info_provided_driver": d.get("bank_info_provided"),
        "vendor_id": d.get("vendor_id"),
        "vendor_name": vendor.get("name"),
        "vendor_description": vendor.get("description"),
        "vendor_is_active": vendor.get("is_active"),
        "assignment_allowed": d.get("assignment_allowed"),
        "rating": d.get("rating"),
        "created_at": d.get("created_at"),
        "updated_at": d.get("updated_at"),
        "hubs": record.get("hubs"),
        "businesses": record.get("businesses"),
        "registered_at": record.get("registered_at"),
        "bank_info_provided": record.get("bank_info_provided"),
        "profile_attendance_status": attendance_latest.get("status"),
        "profile_attendance_lat": attendance_latest.get("lat"),
        "profile_attendance_lon": attendance_latest.get("lon"),
        "profile_attendance_timestamp": attendance_latest.get("time_stamp"),
        "current_lat": coords.get("lat"),
        "current_lon": coords.get("lon"),
        "doc_selfie_photo": doc_selfie.get("fields", {}).get("value", {}).get("photo"),
        "doc_selfie_status": doc_selfie.get("status"),
        "doc_ktp_photo": doc_ktp.get("fields", {}).get("value", {}).get("photo"),
        "doc_ktp_nik": doc_ktp.get("fields", {}).get("value", {}).get("nik"),
        "doc_ktp_status": doc_ktp.get("status"),
        "doc_sim_photo": doc_sim.get("fields", {}).get("value", {}).get("photo"),
        "doc_sim_number": doc_sim.get("fields", {}).get("value", {}).get("sim"),
        "doc_sim_expiry": doc_sim.get("fields", {}).get("value", {}).get("expiry_date"),
        "doc_sim_status": doc_sim.get("status"),
        "profile_bank_name": bank_profile.get("bank_name"),
        "profile_bank_account": bank_profile.get("account"),
        "profile_bank_holder": bank_profile.get("account_holder_name"),
        "appointment_booking_allowed": dp.get("appointment_booking_allowed"),
        "driver_hubs": str(driver_hubs) if driver_hubs else "",
        "driver_businesses": str(driver_businesses) if driver_businesses else "",
        "bank_user_id": bank.get("user_id"),
        "bank_name": bank.get("bank"),
        "bank_account_number": bank.get("account_number"),
        "bank_beneficiary_name": bank.get("beneficiary_name"),
    }


def fetch_all_details(token: str, driver_records: list, max_workers: int = 20) -> dict:
    total = len(driver_records)
    results = {}
    counter = {"done": 0}

    driver_ids = [r.get("drivers", {}).get("id") for r in driver_records]

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(fetch_driver_detail, token, did): did for did in driver_ids if did}
        for future in as_completed(futures):
            driver_id, profile, bank = future.result()
            results[driver_id] = (profile, bank)
            counter["done"] += 1
            if counter["done"] % 100 == 0 or counter["done"] == total:
                safe_print(f"  Detail diambil: {counter['done']}/{total}...")

    return results


def save_to_excel(driver_records: list, detail_map: dict, filename: str):
    if not driver_records:
        print("[WARN] Tidak ada data driver untuk disimpan.")
        return

    rows = []
    for record in driver_records:
        driver_id = record.get("drivers", {}).get("id")
        profile, bank = detail_map.get(driver_id, ({}, {}))
        rows.append(flatten_driver(record, profile, bank))

    columns = list(rows[0].keys())

    wb = openpyxl.Workbook(write_only=False)
    ws = wb.active
    ws.title = "Driver List"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", start_color="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").upper())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    alt_fill = PatternFill("solid", start_color="D6E4F0")
    data_font = Font(name="Arial", size=10)
    data_align = Alignment(vertical="center")

    for row_idx, row in enumerate(rows, start=2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, key in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[key])
            cell.font = data_font
            cell.alignment = data_align
            if fill:
                cell.fill = fill

    ws.row_dimensions[1].height = 30
    for col_idx, col_name in enumerate(columns, start=1):
        values = [len(col_name)] + [len(str(r[col_name])) for r in rows if r[col_name] is not None]
        max_len = max(values) if values else 10
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)
    print(f"[OK] Data disimpan ke '{filename}' ({len(rows)} baris, {len(columns)} kolom)")


def main():
    username = "merapi"
    password = "qRKzNbami4"

    try:
        token = login(username, password)
    except Exception as e:
        print(f"[ERROR] Login gagal: {e}")
        return

    try:
        drivers = fetch_all_drivers(token, offset=100)
    except Exception as e:
        print(f"[ERROR] Gagal mengambil data driver: {e}")
        return

    print(f"[INFO] Mengambil detail profil & bank untuk {len(drivers)} driver (concurrent)...")
    detail_map = fetch_all_details(token, drivers, max_workers=20)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = f"driver_list_{timestamp}.xlsx"
    save_to_excel(drivers, detail_map, output_file)


if __name__ == "__main__":
    main()